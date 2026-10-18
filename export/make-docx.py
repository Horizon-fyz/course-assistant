# 把抓到的试卷生成可打印的 Word 复习稿
#
# 用法：
#   python make-docx.py                                   # 默认读 export/papers.json
#   python make-docx.py --src X.json --out Y.docx
#   python make-docx.py --account "示例用户（示例大学）"
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
import json, os, re, sys, datetime

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _arg(name, default=None):
    a = sys.argv[1:]
    if name in a:
        i = a.index(name)
        if i + 1 < len(a):
            return a[i + 1]
    return default


SRC = _arg("--src", os.path.join(BASE, "export", "papers.json"))
OUT = _arg("--out", os.path.join(BASE, "export", "智慧树作业题库_打印版.docx"))
ACCOUNT = _arg("--account", "")

with open(SRC, encoding="utf-8") as f:
    data = json.load(f)
if not isinstance(data, dict):
    raise SystemExit("papers.json 结构异常：顶层应为 {课程: [作业...]}")

LABELS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def clean(text):
    """去掉接口文本里的 HTML 标签，换行折叠为空格"""
    s = str(text or "")
    s = re.sub(r"<br\s*/?>", " ", s, flags=re.I)
    s = re.sub(r"</(p|div|li)>", " ", s, flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    s = (s.replace("&nbsp;", " ").replace("&lt;", "<").replace("&gt;", ">")
          .replace("&quot;", '"').replace("&amp;", "&"))
    return re.sub(r"\s+", " ", s).strip()


def set_cjk(run, name="宋体", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


doc = Document()

# 页面：A4 + 合理页边距
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for attr, val in (("left_margin", 2.2), ("right_margin", 2.2), ("top_margin", 2.0), ("bottom_margin", 2.0)):
    setattr(sec, attr, Cm(val))

# 正文默认字体
style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

# ---------- 封面 ----------
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cjk(t.add_run("智慧树作业题库"), "微软雅黑", 26, True)

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cjk(s.add_run("复习与打印稿"), "微软雅黑", 14, False, (0x55, 0x55, 0x55))

doc.add_paragraph()

def paper_questions(item):
    """从一份作业里取出全部题目"""
    base = ((item.get("paper") or {}).get("rt") or {}).get("examBase") or {}
    out = []
    for part in base.get("workExamParts") or []:
        out.extend(part.get("questionDtos") or [])
    return out


# 统计
stat = doc.add_paragraph()
stat.alignment = WD_ALIGN_PARAGRAPH.CENTER
total_hw = sum(len(v) for v in data.values() if isinstance(v, list))
total_q = sum(len(paper_questions(x)) for v in data.values() if isinstance(v, list) for x in v)
set_cjk(stat.add_run(f"共 {len(data)} 门课程 · {total_hw} 份作业 · {total_q} 道题目"), "微软雅黑", 12)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cjk(meta.add_run(f"生成时间：{datetime.datetime.now():%Y-%m-%d %H:%M}" + (f"    账号：{ACCOUNT}" if ACCOUNT else "")), "宋体", 9, False, (0x77, 0x77, 0x77))

warn = doc.add_paragraph()
warn.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_cjk(warn.add_run("本稿由已登录会话读取试卷内容生成，仅供本人复习使用；未作答、未提交任何作业。"), "宋体", 9, False, (0x99, 0x33, 0x33))

doc.add_paragraph()

# 目录
doc.add_paragraph()
h = doc.add_paragraph()
set_cjk(h.add_run("目录"), "微软雅黑", 15, True)
for course, items in data.items():
    if not isinstance(items, list):
        continue
    p = doc.add_paragraph()
    set_cjk(p.add_run(f"● {course}（{len(items)} 份作业）"), "微软雅黑", 11, True)
    for it in items:
        base = ((it.get("paper") or {}).get("rt") or {}).get("examBase") or {}
        n = len(paper_questions(it))
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.8)
        set_cjk(q.add_run(f"  {it.get('examName','(未知)')}　{n} 题　{it.get('totalScore','')} 分　截止 {str(it.get('endDate',''))[:10]}"), "宋体", 10)

doc.add_page_break()

# ---------- 正文 ----------
courses = [(c, v) for c, v in data.items() if isinstance(v, list)]
for ci, (course, items) in enumerate(courses):
    ch = doc.add_paragraph()
    set_cjk(ch.add_run(course), "微软雅黑", 20, True, (0x1F, 0x3A, 0x93))
    ch.paragraph_format.space_after = Pt(10)

    for it in items:
        base = ((it.get("paper") or {}).get("rt") or {}).get("examBase") or {}
        parts = base.get("workExamParts") or []
        questions = [q for part in parts for q in (part.get("questionDtos") or [])]
        if not questions:
            continue

        # 作业标题
        hd = doc.add_paragraph()
        hd.paragraph_format.space_before = Pt(14)
        hd.paragraph_format.space_after = Pt(4)
        set_cjk(hd.add_run(f"{it.get('examName','')}"), "微软雅黑", 14, True)

        info = doc.add_paragraph()
        info.paragraph_format.space_after = Pt(8)
        set_cjk(
            info.add_run(
                f"题目数：{it.get('problemNum','')}　总分：{it.get('totalScore','')} 分　"
                f"对应章节：{it.get('cpOrderNumber','')}　截止：{it.get('endDate','')}"
            ),
            "宋体", 9, False, (0x66, 0x66, 0x66),
        )

        for qi, q in enumerate(questions, 1):
            qt = (q.get("questionType") or {}).get("name", "")
            # 题干
            stem = doc.add_paragraph()
            stem.paragraph_format.space_before = Pt(6)
            stem.paragraph_format.space_after = Pt(2)
            set_cjk(stem.add_run(f"{qi}. 【{qt}】 "), "宋体", 10.5, True)
            set_cjk(stem.add_run(clean(q.get("name", ""))), "宋体", 10.5)

            # 选项（保留原始顺序，并明确提示顺序随机）
            for oi, opt in enumerate(q.get("questionOptions") or []):
                op = doc.add_paragraph()
                op.paragraph_format.left_indent = Cm(0.9)
                op.paragraph_format.space_after = Pt(1)
                label = LABELS[oi] if oi < len(LABELS) else str(oi + 1)
                set_cjk(op.add_run(f"{label}. "), "宋体", 10.5, True)
                set_cjk(op.add_run(clean(opt.get("content", ""))), "宋体", 10.5)

        note = doc.add_paragraph()
        note.paragraph_format.space_before = Pt(4)
        set_cjk(note.add_run("（提示：平台选择题选项为随机排列，核对时请以选项内容为准。）"), "宋体", 8.5, False, (0x99, 0x99, 0x99))

        doc.add_paragraph()

    if ci < len(courses) - 1:
        doc.add_page_break()

doc.save(OUT)
print("saved:", OUT)
print("courses:", list(data.keys()))
print("total homework:", total_hw, "total questions:", total_q)
