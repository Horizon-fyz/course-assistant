from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import json, os, sys, datetime

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _arg(name, default=None):
    a = sys.argv[1:]
    if name in a:
        i = a.index(name)
        if i + 1 < len(a):
            return a[i + 1]
    return default


SRC = _arg("--src", os.path.join(BASE, "export", "raw.json"))
OUT = _arg("--out", os.path.join(BASE, "export", "智慧树课程数据.xlsx"))
ACCOUNT = _arg("--account", "")
SCHOOL = _arg("--school", "")

with open(SRC, encoding="utf-8") as f:
    raw = json.load(f)
if not isinstance(raw, dict):
    raise SystemExit("raw.json 结构异常：顶层应为对象")

courses = [c for c in (raw.get("courses") or []) if isinstance(c, dict)]
videos = raw.get("videos") or {}
if not isinstance(videos, dict):
    videos = {}

wb = Workbook()

HDR_FILL = PatternFill("solid", fgColor="2F5597")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=13)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def autofit(ws, minw=8, maxw=52):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            longest = max(longest, sum(2 if ord(ch) > 127 else 1 for ch in v))
        ws.column_dimensions[letter].width = max(minw, min(maxw, longest + 3))


# ---------------- Sheet 1: 课程总览 ----------------
ws = wb.active
ws.title = "课程总览"
ws.append(["课程名称", "教师", "开课学校", "学习进度", "当前在学",
           "courseId", "recruitId", "章节数", "视频数", "已学视频", "视频总时长(分)"])
style_header(ws)

total_chapters = 0
total_lessons = 0
for c in courses:
    v = videos.get(str(c.get("courseId"))) or {}
    chapters = v.get("videoChapterDtos") or []
    lessons = [l for ch in chapters for l in (ch.get("videoLessons") or [])]
    total_chapters += len(chapters)
    total_lessons += len(lessons)
    studied = sum(1 for l in lessons if l.get("isStudiedLesson") == 1)
    total_sec = sum(l.get("videoSec") or 0 for l in lessons)
    ws.append([
        c.get("courseName", ""), c.get("teacherName", ""), c.get("schoolName", ""), c.get("progress", ""),
        f'{c.get("lessonNum") or ""} {c.get("lessonName") or ""}'.strip(),
        c.get("courseId", ""), c.get("recruitId", ""), len(chapters), len(lessons), studied,
        round(total_sec / 60),
    ])
for r in range(2, ws.max_row + 1):
    for cc in range(1, ws.max_column + 1):
        ws.cell(row=r, column=cc).border = BORDER
autofit(ws)
ws.freeze_panes = "A2"

# ---------------- Sheet 2: 章节视频明细 ----------------
ws2 = wb.create_sheet("章节视频明细")
ws2.append(["课程名称", "章节序号", "章节名", "课节序号", "课节名称",
            "时长", "时长(秒)", "是否已学", "lessonId", "videoId", "授课教师"])
style_header(ws2)

for c in courses:
    v = videos.get(str(c.get("courseId"))) or {}
    for ch in v.get("videoChapterDtos") or []:
        for l in ch.get("videoLessons") or []:
            sec = l.get("videoSec") or 0
            ws2.append([
                c.get("courseName", ""), ch.get("orderNumber"), ch.get("name"),
                l.get("orderNumber"), l.get("name"),
                f"{sec // 60:02d}:{sec % 60:02d}", sec,
                "已学" if l.get("isStudiedLesson") == 1 else "未学",
                l.get("id"), l.get("videoId"), l.get("teacherName"),
            ])
for r in range(2, ws2.max_row + 1):
    for cc in range(1, ws2.max_column + 1):
        ws2.cell(row=r, column=cc).border = BORDER
autofit(ws2)
ws2.freeze_panes = "A2"

# ---------------- Sheet 3: 数据来源说明 ----------------
ws3 = wb.create_sheet("数据来源与说明")
ws3["A1"] = "数据来源与说明"
ws3["A1"].font = TITLE_FONT
rows = [
    ("采集时间", raw.get("fetchedAt")),
    ("账号", ACCOUNT),
    ("学校", SCHOOL),
    ("", ""),
    ("课程列表接口", "POST onlineservice-api.zhihuishu.com/gateway/t/v1/student/course/share/queryShareCourseInfo"),
    ("章节视频接口", "POST studyservice-api.zhihuishu.com/gateway/t/v1/learning/videolist"),
    ("", ""),
    ("已采集", f"课程 {len(courses)} 门、章节 {total_chapters} 个、视频 {total_lessons} 个（含时长与已学标记）"),
    ("未采集", "单课节观看时长（queryStuyInfo）"),
    ("未采集原因", "该接口要求把 recruitId/courseId 加密进 secretStr，无法与其他接口复用签名，"
                   "服务端返回 code 421「该账号无权限」；需通过浏览器实际操作该页面采集。"),
    ("", ""),
    ("说明", "学习进度取自课程列表接口返回的 progress 字段，为平台计算值。"),
]
r = 3
for k, v in rows:
    ws3.cell(row=r, column=1, value=k).font = Font(bold=bool(k))
    ws3.cell(row=r, column=2, value=v)
    ws3.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
ws3.column_dimensions["A"].width = 16
ws3.column_dimensions["B"].width = 96

wb.save(OUT)
print("saved:", OUT)
print("sheets:", wb.sheetnames)
print("课程总览 rows:", ws.max_row - 1, "| 明细 rows:", ws2.max_row - 1)
